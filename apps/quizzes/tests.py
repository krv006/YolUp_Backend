import re
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.accounts.tests import register, login

from .models import Quiz


def _build_docx(lines: list) -> bytes:
    import docx

    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_xlsx(rows: list) -> bytes:
    """`rows` — har biri `[savol, variantA, variantB, ..., javob_harfi]` ro'yxati."""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['Savol', 'A', 'B', 'C', 'D', 'E', "To'g'ri javob"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _mcq(text, correct_index, options, points=1):
    return {
        'text': text,
        'points': points,
        'options': [
            {'text': opt, 'is_correct': i == correct_index}
            for i, opt in enumerate(options)
        ],
    }


class QuizFlowTests(APITestCase):
    def setUp(self):
        register(self.client, 't1', 'teacher')
        self.teacher_token = login(self.client, 't1')

        register(self.client, 'p1', 'parent')
        self.parent_token = login(self.client, 'p1')
        self.auth(self.parent_token)
        resp = self.client.post('/api/v1/auth/children/', {'username': 's1', 'password': 'StrongPass123!'})
        self.child_id = resp.json()['id']
        self.child_token = login(self.client, 's1')

        self.auth(self.teacher_token)
        self.course_id = self.client.post(
            '/api/v1/courses/', {'title': 'Algebra', 'subject': 'Matematika'}
        ).json()['id']

        self.quiz_payload = {
            'course': self.course_id,
            'title': "1-bob testi",
            'questions': [
                _mcq('2 + 2 = ?', correct_index=1, options=['3', '4', '5']),
                _mcq('Poytaxt?', correct_index=0, options=['Toshkent', 'Samarqand'], points=2),
            ],
        }

    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def enroll_child(self):
        self.auth(self.parent_token)
        self.client.post(f'/api/v1/courses/{self.course_id}/enroll/', {'student_id': self.child_id})
        self.auth(self.teacher_token)
        for req in self.client.get('/api/v1/courses/requests/').json()['results']:
            self.client.post('/api/v1/courses/requests/respond/', {
                'enrollment_id': req['id'], 'action': 'approve',
            })

    def create_quiz(self):
        self.auth(self.teacher_token)
        resp = self.client.post('/api/v1/quizzes/', self.quiz_payload, format='json')
        return resp

    def test_teacher_creates_quiz_with_questions_and_options(self):
        resp = self.create_quiz()
        self.assertEqual(resp.status_code, 201)
        self.assertEqual(len(resp.json()['questions']), 2)
        self.assertEqual(len(resp.json()['questions'][0]['options']), 3)
        self.assertTrue(Quiz.objects.filter(title="1-bob testi").exists())

    def test_student_cannot_create_quiz(self):
        self.enroll_child()
        self.auth(self.child_token)
        resp = self.client.post('/api/v1/quizzes/', self.quiz_payload, format='json')
        self.assertEqual(resp.status_code, 403)

    def test_question_requires_exactly_one_correct_option(self):
        self.auth(self.teacher_token)
        bad = {
            'course': self.course_id,
            'title': 'Xato test',
            'questions': [{
                'text': '1+1=?', 'points': 1,
                'options': [{'text': '2', 'is_correct': True}, {'text': '3', 'is_correct': True}],
            }],
        }
        resp = self.client.post('/api/v1/quizzes/', bad, format='json')
        self.assertEqual(resp.status_code, 400)

    def test_student_take_view_hides_correct_answer(self):
        quiz_id = self.create_quiz().json()['id']
        self.enroll_child()
        self.auth(self.child_token)
        resp = self.client.get(f'/api/v1/quizzes/{quiz_id}/')
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn('is_correct', resp.json()['questions'][0]['options'][0])

    def test_non_enrolled_student_cannot_see_quiz(self):
        quiz_id = self.create_quiz().json()['id']
        self.auth(self.child_token)  # hali enroll qilinmagan
        resp = self.client.get(f'/api/v1/quizzes/{quiz_id}/')
        self.assertEqual(resp.status_code, 404)

    def test_submit_attempt_scores_correctly_and_reveals_answers(self):
        quiz = self.create_quiz().json()
        self.enroll_child()
        q1, q2 = quiz['questions']
        answers = [
            {'question': q1['id'], 'selected_option': q1['options'][1]['id']},  # to'g'ri (4)
            {'question': q2['id'], 'selected_option': q2['options'][1]['id']},  # xato (Samarqand)
        ]
        self.auth(self.child_token)
        resp = self.client.post(f'/api/v1/quizzes/{quiz["id"]}/attempts/', {'answers': answers}, format='json')
        self.assertEqual(resp.status_code, 201)
        body = resp.json()
        self.assertEqual(body['score'], 1)
        self.assertEqual(body['max_score'], 3)
        wrong_answer = next(a for a in body['answers'] if not a['is_correct'])
        self.assertEqual(wrong_answer['correct_option']['text'], 'Toshkent')

    def test_student_can_retake_unlimited_times(self):
        quiz = self.create_quiz().json()
        self.enroll_child()
        q1, q2 = quiz['questions']
        self.auth(self.child_token)
        for _ in range(3):
            self.client.post(f'/api/v1/quizzes/{quiz["id"]}/attempts/', {'answers': [
                {'question': q1['id'], 'selected_option': q1['options'][1]['id']},
                {'question': q2['id'], 'selected_option': q2['options'][0]['id']},
            ]}, format='json')
        resp = self.client.get(f'/api/v1/quizzes/{quiz["id"]}/attempts/')
        self.assertEqual(len(resp.json()), 3)
        self.assertTrue(all(a['score'] == 3 for a in resp.json()))

    def test_other_teacher_cannot_see_or_delete_foreign_quiz(self):
        quiz_id = self.create_quiz().json()['id']
        register(self.client, 't2', 'teacher')
        self.auth(login(self.client, 't2'))
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{quiz_id}/').status_code, 404)
        self.assertEqual(self.client.delete(f'/api/v1/quizzes/{quiz_id}/').status_code, 404)

    def test_owning_teacher_can_delete_quiz(self):
        quiz_id = self.create_quiz().json()['id']
        self.auth(self.teacher_token)
        resp = self.client.delete(f'/api/v1/quizzes/{quiz_id}/')
        self.assertEqual(resp.status_code, 204)
        self.assertFalse(Quiz.objects.filter(pk=quiz_id).exists())

    def test_creating_quiz_notifies_enrolled_students(self):
        from apps.notifications.models import Notification

        self.enroll_child()
        with self.captureOnCommitCallbacks(execute=True):
            self.create_quiz()
        note = Notification.objects.filter(link_type='quiz').latest('created_at')
        self.assertIn('1-bob testi', note.description)
        self.assertTrue(note.recipients.filter(user_id=self.child_id).exists())

    def test_future_quiz_hidden_from_student_until_opens_at(self):
        from datetime import timedelta

        from django.utils import timezone

        self.enroll_child()
        self.auth(self.teacher_token)
        payload = {**self.quiz_payload, 'opens_at': (timezone.now() + timedelta(days=1)).isoformat()}
        quiz_id = self.client.post('/api/v1/quizzes/', payload, format='json').json()['id']

        self.auth(self.child_token)
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{quiz_id}/').status_code, 404)

        self.auth(self.teacher_token)
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{quiz_id}/').status_code, 200)

    def test_future_quiz_does_not_notify_yet(self):
        from datetime import timedelta

        from django.utils import timezone

        from apps.notifications.models import Notification

        self.enroll_child()
        self.auth(self.teacher_token)
        payload = {**self.quiz_payload, 'opens_at': (timezone.now() + timedelta(days=1)).isoformat()}
        self.client.post('/api/v1/quizzes/', payload, format='json')
        self.assertFalse(Notification.objects.filter(link_type='quiz').exists())


class MockTestFlowTests(APITestCase):
    def setUp(self):
        register(self.client, 'mt1', 'teacher')
        self.teacher_token = login(self.client, 'mt1')

        register(self.client, 'mp1', 'parent')
        self.parent_token = login(self.client, 'mp1')
        self.auth(self.parent_token)
        resp = self.client.post('/api/v1/auth/children/', {'username': 'ms1', 'password': 'StrongPass123!'})
        self.child_id = resp.json()['id']
        self.child_token = login(self.client, 'ms1')

        self.auth(self.teacher_token)
        self.course_id = self.client.post(
            '/api/v1/courses/', {'title': 'Ingliz tili', 'subject': 'Til'}
        ).json()['id']

        self.quiz1_id = self._create_quiz('1-bob testi')['id']
        self.quiz2_id = self._create_quiz('2-bob testi')['id']

    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def _create_quiz(self, title, course_id=None):
        self.auth(self.teacher_token)
        payload = {
            'course': course_id or self.course_id,
            'title': title,
            'questions': [_mcq('2 + 2 = ?', correct_index=1, options=['3', '4', '5'])],
        }
        return self.client.post('/api/v1/quizzes/', payload, format='json').json()

    def enroll_child(self):
        self.auth(self.parent_token)
        self.client.post(f'/api/v1/courses/{self.course_id}/enroll/', {'student_id': self.child_id})
        self.auth(self.teacher_token)
        for req in self.client.get('/api/v1/courses/requests/').json()['results']:
            self.client.post('/api/v1/courses/requests/respond/', {
                'enrollment_id': req['id'], 'action': 'approve',
            })

    def create_mock_test(self, quiz_ids=None):
        self.auth(self.teacher_token)
        payload = {
            'course': self.course_id,
            'title': "1-chorak yakuniy imtihoni",
            'time_limit_minutes': 30,
            'quizzes': quiz_ids or [self.quiz1_id, self.quiz2_id],
        }
        return self.client.post('/api/v1/quizzes/mock-tests/', payload, format='json')

    def test_teacher_creates_mock_test_from_existing_quizzes(self):
        resp = self.create_mock_test()
        self.assertEqual(resp.status_code, 201)
        body = resp.json()
        self.assertEqual(len(body['sections']), 2)
        self.assertEqual(body['sections'][0]['quiz_id'], self.quiz1_id)
        self.assertEqual(body['sections'][1]['quiz_id'], self.quiz2_id)

    def test_student_cannot_create_mock_test(self):
        self.enroll_child()
        self.auth(self.child_token)
        resp = self.client.post('/api/v1/quizzes/mock-tests/', {
            'course': self.course_id, 'title': 'X', 'time_limit_minutes': 10, 'quizzes': [self.quiz1_id],
        }, format='json')
        self.assertEqual(resp.status_code, 403)

    def test_rejects_quiz_from_other_course(self):
        other_course_id = self.client.post(
            '/api/v1/courses/', {'title': 'Boshqa kurs', 'subject': 'X'}
        ).json()['id']
        foreign_quiz_id = self._create_quiz('Boshqa test', course_id=other_course_id)['id']
        resp = self.create_mock_test(quiz_ids=[self.quiz1_id, foreign_quiz_id])
        self.assertEqual(resp.status_code, 400)

    def test_student_starts_mock_test_and_sees_questions_without_answers(self):
        mock_test_id = self.create_mock_test().json()['id']
        self.enroll_child()
        self.auth(self.child_token)
        resp = self.client.post(f'/api/v1/quizzes/mock-tests/{mock_test_id}/start/')
        self.assertEqual(resp.status_code, 201)
        body = resp.json()
        self.assertIn('deadline', body)
        self.assertEqual(len(body['sections']), 2)
        first_option = body['sections'][0]['quiz']['questions'][0]['options'][0]
        self.assertNotIn('is_correct', first_option)

    def test_non_enrolled_student_cannot_start(self):
        mock_test_id = self.create_mock_test().json()['id']
        self.auth(self.child_token)  # hali enroll qilinmagan
        resp = self.client.post(f'/api/v1/quizzes/mock-tests/{mock_test_id}/start/')
        self.assertEqual(resp.status_code, 404)

    def test_submit_scores_all_sections_and_creates_quiz_attempts(self):
        from .models import QuizAttempt

        mock_test = self.create_mock_test().json()
        self.enroll_child()
        self.auth(self.child_token)
        start = self.client.post(f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/start/').json()
        q1 = start['sections'][0]['quiz']
        q2 = start['sections'][1]['quiz']

        payload = {'sections': [
            {'quiz': q1['id'], 'answers': [
                {'question': q1['questions'][0]['id'], 'selected_option': q1['questions'][0]['options'][1]['id']},
            ]},
            {'quiz': q2['id'], 'answers': [
                {'question': q2['questions'][0]['id'], 'selected_option': q2['questions'][0]['options'][0]['id']},
            ]},
        ]}
        resp = self.client.post(
            f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/attempts/{start["id"]}/submit/',
            payload, format='json',
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertEqual(body['total_score'], 1)  # faqat q1 to'g'ri
        self.assertEqual(body['total_max_score'], 2)
        self.assertEqual(len(body['quiz_results']), 2)
        self.assertEqual(
            QuizAttempt.objects.filter(mock_test_attempt_id=start['id']).count(), 2,
        )

    def test_cannot_submit_twice(self):
        mock_test = self.create_mock_test().json()
        self.enroll_child()
        self.auth(self.child_token)
        start = self.client.post(f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/start/').json()
        payload = {'sections': [
            {'quiz': s['quiz']['id'], 'answers': [
                {'question': s['quiz']['questions'][0]['id'],
                 'selected_option': s['quiz']['questions'][0]['options'][0]['id']},
            ]}
            for s in start['sections']
        ]}
        submit_url = f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/attempts/{start["id"]}/submit/'
        self.assertEqual(self.client.post(submit_url, payload, format='json').status_code, 200)
        self.assertEqual(self.client.post(submit_url, payload, format='json').status_code, 400)

    def test_cannot_submit_after_deadline(self):
        from datetime import timedelta

        from django.utils import timezone

        from .models import MockTestAttempt

        mock_test = self.create_mock_test().json()
        self.enroll_child()
        self.auth(self.child_token)
        start = self.client.post(f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/start/').json()
        # `auto_now_add` — `.update()` orqali chetlab o'tib, muddatni "o'tgan" qilamiz.
        MockTestAttempt.objects.filter(pk=start['id']).update(
            started_at=timezone.now() - timedelta(minutes=60),
        )
        payload = {'sections': [
            {'quiz': s['quiz']['id'], 'answers': [
                {'question': s['quiz']['questions'][0]['id'],
                 'selected_option': s['quiz']['questions'][0]['options'][0]['id']},
            ]}
            for s in start['sections']
        ]}
        resp = self.client.post(
            f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/attempts/{start["id"]}/submit/',
            payload, format='json',
        )
        self.assertEqual(resp.status_code, 400)

    def test_teacher_sees_all_attempts_student_sees_only_own(self):
        mock_test = self.create_mock_test().json()
        self.enroll_child()
        self.auth(self.child_token)
        self.client.post(f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/start/')

        self.auth(self.teacher_token)
        resp = self.client.get(f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/attempts/')
        self.assertEqual(len(resp.json()), 1)

        self.auth(self.child_token)
        resp = self.client.get(f'/api/v1/quizzes/mock-tests/{mock_test["id"]}/attempts/')
        self.assertEqual(len(resp.json()), 1)


class QuizDocxImportTests(APITestCase):
    def setUp(self):
        register(self.client, 't1', 'teacher')
        self.teacher_token = login(self.client, 't1')
        register(self.client, 's1', 'student')
        self.student_token = login(self.client, 's1')

    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def import_docx(self, lines):
        self.auth(self.teacher_token)
        content = _build_docx(lines)
        upload = SimpleUploadedFile(
            'test.docx', content,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        return self.client.post('/api/v1/quizzes/import/', {'file': upload}, format='multipart')

    def test_teacher_imports_well_formed_docx(self):
        resp = self.import_docx([
            'Nevrologiya fanidan test savollari',
            'Bolalar serebral falaji mavzusi',
            "1. Qaysi shaklda mushak tonusi oshadi?",
            'A) Giperkinetik shakl',
            'B) Spastik diplegiya',
            'C) Miyachali shakl',
            "To'g'ri javob: B",
            '2. Ataksiya qaysi shaklga xos?',
            'A) Miyachali shakl',
            'B) Spastik shakl',
            "To'g'ri javob: A",
        ])
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data['title'], 'Nevrologiya fanidan test savollari')
        self.assertEqual(data['description'], 'Bolalar serebral falaji mavzusi')
        self.assertEqual(len(data['questions']), 2)
        self.assertEqual(data['warnings'], [])
        q1_options = data['questions'][0]['options']
        self.assertEqual(len(q1_options), 3)
        self.assertEqual(sum(1 for o in q1_options if o['is_correct']), 1)
        self.assertTrue(q1_options[1]['is_correct'])  # B = index 1

    def test_import_flags_question_with_no_detected_answer(self):
        resp = self.import_docx([
            '1. Javobi yo\'q savol',
            'A) Variant 1',
            'B) Variant 2',
        ])
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(len(data['warnings']), 1)
        self.assertEqual(data['warnings'][0]['reason'], 'answer_not_detected')
        self.assertEqual(sum(1 for o in data['questions'][0]['options'] if o['is_correct']), 0)

    def test_import_rejects_file_with_no_questions(self):
        resp = self.import_docx(['Bu yerda birorta ham savol yo\'q'])
        self.assertEqual(resp.status_code, 400)

    def test_import_rejects_non_docx_extension(self):
        self.auth(self.teacher_token)
        upload = SimpleUploadedFile('test.txt', b'1. Savol?\nA) X\nB) Y\nJavob: A', content_type='text/plain')
        resp = self.client.post('/api/v1/quizzes/import/', {'file': upload}, format='multipart')
        self.assertEqual(resp.status_code, 400)

    def test_student_cannot_import(self):
        self.auth(self.student_token)
        content = _build_docx(['1. Savol?', 'A) X', 'B) Y', "To'g'ri javob: A"])
        upload = SimpleUploadedFile('test.docx', content)
        resp = self.client.post('/api/v1/quizzes/import/', {'file': upload}, format='multipart')
        self.assertEqual(resp.status_code, 403)

    def test_imported_preview_can_be_submitted_as_real_quiz(self):
        resp = self.import_docx([
            '1. 2+2 nechchi?',
            'A) 3',
            'B) 4',
            "To'g'ri javob: B",
        ])
        preview = resp.json()
        self.auth(self.teacher_token)
        course_id = self.client.post(
            '/api/v1/courses/', {'title': 'Matematika', 'subject': 'Matematika'}
        ).json()['id']
        create_resp = self.client.post('/api/v1/quizzes/', {
            'course': course_id,
            'title': preview['title'] or 'Import qilingan test',
            'description': preview['description'],
            'questions': preview['questions'],
        }, format='json')
        self.assertEqual(create_resp.status_code, 201)


class QuizXlsxImportTests(APITestCase):
    def setUp(self):
        register(self.client, 'xt1', 'teacher')
        self.teacher_token = login(self.client, 'xt1')
        register(self.client, 'xs1', 'student')
        self.student_token = login(self.client, 'xs1')

    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def import_xlsx(self, rows):
        self.auth(self.teacher_token)
        content = _build_xlsx(rows)
        upload = SimpleUploadedFile(
            'test.xlsx', content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        return self.client.post('/api/v1/quizzes/import/', {'file': upload}, format='multipart')

    def test_teacher_imports_well_formed_xlsx(self):
        resp = self.import_xlsx([
            ['Poytaxt qaysi shahar?', 'Samarqand', 'Toshkent', 'Buxoro', '', '', 'B'],
            ['2+2 nechchi?', '3', '4', '', '', '', 'B'],
        ])
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(len(data['questions']), 2)
        self.assertEqual(data['warnings'], [])
        q1_options = data['questions'][0]['options']
        self.assertEqual(len(q1_options), 3)
        self.assertTrue(q1_options[1]['is_correct'])  # B = index 1 (Toshkent)

    def test_xlsx_skips_blank_rows(self):
        resp = self.import_xlsx([
            ['Savol 1?', 'X', 'Y', '', '', '', 'A'],
            ['', '', '', '', '', '', ''],
            ['Savol 2?', 'X', 'Y', '', '', '', 'B'],
        ])
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.json()['questions']), 2)

    def test_xlsx_flags_question_with_no_detected_answer(self):
        resp = self.import_xlsx([
            ['Javobsiz savol', 'X', 'Y', '', '', '', ''],
        ])
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(len(data['warnings']), 1)
        self.assertEqual(data['warnings'][0]['reason'], 'answer_not_detected')

    def test_xlsx_rejects_file_with_no_questions(self):
        resp = self.import_xlsx([])
        self.assertEqual(resp.status_code, 400)

    def test_student_cannot_import_xlsx(self):
        self.auth(self.student_token)
        content = _build_xlsx([['Savol?', 'X', 'Y', '', '', '', 'A']])
        upload = SimpleUploadedFile('test.xlsx', content)
        resp = self.client.post('/api/v1/quizzes/import/', {'file': upload}, format='multipart')
        self.assertEqual(resp.status_code, 403)

    def test_imported_xlsx_preview_can_be_submitted_as_real_quiz(self):
        resp = self.import_xlsx([['2+2 nechchi?', '3', '4', '', '', '', 'B']])
        preview = resp.json()
        self.auth(self.teacher_token)
        course_id = self.client.post(
            '/api/v1/courses/', {'title': 'Matematika', 'subject': 'Matematika'}
        ).json()['id']
        create_resp = self.client.post('/api/v1/quizzes/', {
            'course': course_id,
            'title': 'Xlsx testi',
            'questions': preview['questions'],
        }, format='json')
        self.assertEqual(create_resp.status_code, 201)


class QuizTemplateExportTests(APITestCase):
    """Shablonni yuklab olish (`GET /quizzes/template/`) — va ENG MUHIMI,
    to'ldirilgan shablon qaytadan import qilinganda muammosiz o'tishi."""

    def setUp(self):
        register(self.client, 'tt1', 'teacher')
        self.teacher_token = login(self.client, 'tt1')
        register(self.client, 'ts1', 'student')
        self.student_token = login(self.client, 'ts1')
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.teacher_token}')

    def test_download_docx_template(self):
        resp = self.client.get('/api/v1/quizzes/template/?type=docx&count=3')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        self.assertIn('attachment', resp['Content-Disposition'])

    def test_download_xlsx_template(self):
        resp = self.client.get('/api/v1/quizzes/template/?type=xlsx&count=3')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment', resp['Content-Disposition'])

    def test_template_respects_requested_question_count(self):
        from apps.quizzes import docx_import

        resp = self.client.get('/api/v1/quizzes/template/?type=docx&count=5')
        result = docx_import.parse_docx_questions(BytesIO(resp.content))
        self.assertEqual(len(result['questions']), 5)

    def test_template_count_is_clamped(self):
        from apps.quizzes import docx_import

        resp = self.client.get('/api/v1/quizzes/template/?type=docx&count=99999')
        result = docx_import.parse_docx_questions(BytesIO(resp.content))
        self.assertEqual(len(result['questions']), 100)  # MAX_TEMPLATE_QUESTIONS

    def test_student_cannot_download_template(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.student_token}')
        resp = self.client.get('/api/v1/quizzes/template/?type=docx&count=3')
        self.assertEqual(resp.status_code, 403)

    def test_filled_docx_template_round_trips_without_warnings(self):
        """Shablonni yuklab olib, har bir maydonni oddiy o'quvchi kabi
        to'ldirilgach — qayta import qilganda birorta ham warning
        bo'lmasligi kerak (asosiy talab: import muammosiz o'tishi)."""
        import docx

        from apps.quizzes import docx_import

        resp = self.client.get('/api/v1/quizzes/template/?type=docx&count=2')
        document = docx.Document(BytesIO(resp.content))

        filled = docx.Document()
        for paragraph in document.paragraphs:
            text = paragraph.text
            if 'Savol matnini shu yerga yozing' in text:
                text = text.replace('[Savol matnini shu yerga yozing]', "Poytaxt qaysi shahar?")
            elif '-variant matni' in text:
                text = re.sub(r'\[[A-E]-variant matni\]', 'Toshkent', text)
            elif "A/B/C/D dan birini yozing" in text:
                text = "To'g'ri javob: A"
            filled.add_paragraph(text)
        buffer = BytesIO()
        filled.save(buffer)
        buffer.seek(0)

        result = docx_import.parse_docx_questions(buffer)
        self.assertEqual(len(result['questions']), 2)
        self.assertEqual(result['warnings'], [])
        for question in result['questions']:
            self.assertEqual(sum(1 for o in question['options'] if o['is_correct']), 1)

    def test_filled_xlsx_template_round_trips_without_warnings(self):
        import openpyxl

        from apps.quizzes import xlsx_import

        resp = self.client.get('/api/v1/quizzes/template/?type=xlsx&count=2')
        workbook = openpyxl.load_workbook(BytesIO(resp.content))
        sheet = workbook.active
        # Ustunlar soni doim fiks: Savol | A | B | C | D | E | To'g'ri javob — 7 ustun (0-6 indeks).
        self.assertEqual(sheet.max_column, 7)
        for row in sheet.iter_rows(min_row=2):
            row[0].value = 'Poytaxt qaysi shahar?'
            row[1].value = 'Samarqand'
            row[2].value = 'Toshkent'
            row[3].value = 'Buxoro'
            row[4].value = 'Andijon'
            row[5].value = None
            row[6].value = 'B'
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        result = xlsx_import.parse_xlsx_questions(buffer)
        self.assertEqual(len(result['questions']), 2)
        self.assertEqual(result['warnings'], [])
        for question in result['questions']:
            self.assertEqual(sum(1 for o in question['options'] if o['is_correct']), 1)
