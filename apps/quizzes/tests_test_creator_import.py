"""Test-creator (RJalol/Test-creator) eksportlarini import qilish: DOCX, XLSX, JSON."""
import io
import json

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase

from . import docx_import, test_creator_import, xlsx_import
from .models import Question, Quiz
from .tests_question_types import QuizTestBase

ROWS = [
    {'order': 1, 'question_text': 'What is the capital of France?',
     'options': [('A', 'Berlin'), ('B', 'Paris'), ('C', 'Rome'), ('D', 'Madrid')],
     'correct': ['B'], 'explanation': 'Paris is the capital.', 'points': 1.0},
    {'order': 2, 'question_text': '2 + 2 = ?',
     'options': [('A', '3'), ('B', '4'), ('C', '5')], 'correct': ['B'], 'explanation': '', 'points': 2.0},
    {'order': 3, 'question_text': 'Which are prime numbers?',
     'options': [('A', '2'), ('B', '4'), ('C', '7'), ('D', '9')], 'correct': ['A', 'C'],
     'explanation': 'Primes: 2, 7. This explanation is long\nand wraps to the next paragraph.', 'points': 3.0},
]


def build_docx(mode='teacher'):
    from docx import Document

    doc = Document()
    doc.add_heading('IELTS Reading mock', level=1)
    doc.add_paragraph(f"Versiya 1 · {'O’qituvchi nusxasi' if mode == 'teacher' else 'Talaba nusxasi'}")
    for row in ROWS:
        doc.add_paragraph(f"{row['order']}. {row['question_text']}", style='List Number')
        for label, text in row['options']:
            doc.add_paragraph(f' {label}) {text}')
        if mode == 'teacher':
            doc.add_paragraph(f" To'g'ri javob: {', '.join(row['correct'])}")
            for index, part in enumerate(p for p in row['explanation'].split('\n') if p):
                # Izoh birinchi qatorda `Izoh:` bilan, davomi (bir necha paragraf bo'lsa) prefikssiz
                doc.add_paragraph(f' Izoh: {part}' if index == 0 else f' {part}')
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def build_xlsx(mode='teacher'):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = 'Test'
    headers = ['#', 'Savol', 'Turi', 'Qiyinlik', 'Variantlar', 'Ball']
    if mode == 'teacher':
        headers += ["To'g'ri javob", 'Izoh']
    ws.append(headers)
    for row in ROWS:
        options = '; '.join(f'{label}) {text}' for label, text in row['options'])
        line = [row['order'], row['question_text'], 'mcq', 'medium', options, row['points']]
        if mode == 'teacher':
            line += [', '.join(row['correct']), row['explanation']]
        ws.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_json(mode='teacher', extra=None):
    questions = []
    for row in ROWS:
        item = {
            'order': row['order'], 'question_text': row['question_text'], 'question_type': 'mcq',
            'difficulty': 'medium', 'points': row['points'],
            'options': [{'label': l, 'text': t} for l, t in row['options']],
        }
        if mode == 'teacher':
            item.update(correct_options=row['correct'], correct_answer={'labels': row['correct']},
                        explanation=row['explanation'])
        questions.append(item)
    questions += extra or []
    return json.dumps({'test_id': 'x', 'title': 'IELTS Reading mock', 'version': 2, 'mode': mode,
                       'questions': questions}, ensure_ascii=False).encode('utf-8')


def correct_texts(question):
    return [o['text'] for o in question['options'] if o['is_correct']]


class DocxExportTests(SimpleTestCase):
    def test_teacher_export_is_parsed_with_answers_and_explanations_dropped(self):
        result = docx_import.parse_docx_questions(build_docx('teacher'))
        self.assertEqual(result['title'], 'IELTS Reading mock')
        self.assertEqual(len(result['questions']), 3)
        first, second, third = result['questions']
        self.assertEqual(correct_texts(first), ['Paris'])
        self.assertEqual([o['text'] for o in first['options']], ['Berlin', 'Paris', 'Rome', 'Madrid'])  # Izoh yopishmadi
        self.assertEqual(correct_texts(second), ['4'])
        self.assertEqual(third['type'], 'multiple')
        self.assertEqual(correct_texts(third), ['2', '7'])
        self.assertEqual([o['text'] for o in third['options']], ['2', '4', '7', '9'])
        self.assertEqual(result['warnings'], [])

    def test_student_export_has_no_answers_and_warns(self):
        result = docx_import.parse_docx_questions(build_docx('student'))
        self.assertEqual(len(result['questions']), 3)
        self.assertEqual({w['reason'] for w in result['warnings']}, {'answer_not_detected'})
        self.assertEqual(len(result['warnings']), 3)

    def test_answer_line_variants(self):
        for line, expected in (("To'g'ri javob: B", ['B']), ('Javob: A, C', ['A', 'C']), ('Answer: A and D', ['A', 'D']),
                               ("To'g'ri javob: a/c.", ['A', 'C']), ('Javob: A va B', ['A', 'B'])):
            result = docx_import.parse_lines(['1. Savol', 'A) a', 'B) b', 'C) c', 'D) d', line])
            letters = [chr(65 + i) for i, o in enumerate(result['questions'][0]['options']) if o['is_correct']]
            self.assertEqual(letters, expected, line)

    def test_answer_line_that_is_not_letters_stays_an_option_continuation(self):
        result = docx_import.parse_lines(['1. Savol', 'A) kecha', 'Javob: yo\'q', 'B) bugun'])
        self.assertEqual(result['questions'][0]['options'][0]['text'], "kecha Javob: yo'q")

    def test_explanation_wrapping_over_several_lines_is_dropped(self):
        result = docx_import.parse_lines([
            '1. Savol', 'A) a', 'B) b', "To'g'ri javob: A", 'Izoh: birinchi qator', 'davomi shu yerda', '2. Ikkinchi', 'A) x', 'B) y'])
        self.assertEqual([len(q['options']) for q in result['questions']], [2, 2])
        self.assertEqual(result['questions'][0]['options'][1]['text'], 'b')


class XlsxExportTests(SimpleTestCase):
    def test_teacher_export(self):
        result = xlsx_import.parse_xlsx_questions(build_xlsx('teacher'))
        self.assertEqual(len(result['questions']), 3)
        first, second, third = result['questions']
        self.assertEqual(correct_texts(first), ['Paris'])
        self.assertEqual(first['points'], 1)
        self.assertEqual(second['points'], 2)
        self.assertEqual(third['type'], 'multiple')
        self.assertEqual(correct_texts(third), ['2', '7'])
        self.assertEqual(third['points'], 3)
        self.assertEqual(result['warnings'], [])

    def test_student_export_warns(self):
        result = xlsx_import.parse_xlsx_questions(build_xlsx('student'))
        self.assertEqual(len(result['questions']), 3)
        self.assertEqual(len(result['warnings']), 3)

    def test_our_own_layout_still_works(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(['Savol', 'A', 'B', 'C', 'D', 'E', "To'g'ri javob"])
        ws.append(['Savol 1', 'bir', 'ikki', 'uch', 'to\'rt', None, 'B'])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        result = xlsx_import.parse_xlsx_questions(buf)
        self.assertEqual(correct_texts(result['questions'][0]), ['ikki'])
        self.assertNotIn('points', result['questions'][0])


class JsonExportTests(SimpleTestCase):
    def parse(self, payload):
        return test_creator_import.parse_test_creator_json(payload)

    def test_teacher_mode(self):
        result = self.parse(build_json('teacher'))
        self.assertEqual(result['title'], 'IELTS Reading mock')
        self.assertIn('versiya 2', result['description'])
        first, second, third = result['questions']
        self.assertEqual(correct_texts(first), ['Paris'])
        self.assertEqual((first['points'], second['points'], third['points']), (1, 2, 3))
        self.assertEqual(third['type'], 'multiple')
        self.assertEqual(correct_texts(third), ['2', '7'])
        self.assertEqual(result['warnings'], [])

    def test_student_mode_has_no_answers(self):
        result = self.parse(build_json('student'))
        self.assertEqual(len(result['warnings']), 3)

    def test_questions_are_sorted_by_order(self):
        payload = json.loads(build_json('teacher'))
        payload['questions'].reverse()
        result = self.parse(json.dumps(payload).encode())
        self.assertEqual(result['questions'][0]['text'], 'What is the capital of France?')

    def test_plain_list_payload_and_bom(self):
        payload = json.loads(build_json('teacher'))['questions']
        result = self.parse(b'\xef\xbb\xbf' + json.dumps(payload).encode())
        self.assertEqual(len(result['questions']), 3)

    def test_non_option_question_types(self):
        extra = [
            {'order': 4, 'question_text': 'The earth is flat.', 'question_type': 'true_false', 'correct_answer': 'false'},
            {'order': 5, 'question_text': '7 x 8 = ?', 'question_type': 'numeric', 'correct_answer': {'value': 56}},
            {'order': 6, 'question_text': 'Capital of Uzbekistan?', 'question_type': 'short_answer',
             'correct_answer': ['Toshkent', 'Tashkent']},
            {'order': 7, 'question_text': 'Javobsiz savol', 'question_type': 'short_answer', 'correct_answer': {}},
        ]
        result = self.parse(build_json('teacher', extra))
        tf, num, text, empty = result['questions'][3:]
        self.assertEqual((tf['type'], tf['correct_bool']), ('true_false', False))
        self.assertEqual((num['type'], num['accepted_answers']), ('numeric', ['56']))
        self.assertEqual((text['type'], text['accepted_answers']), ('text', ['Toshkent', 'Tashkent']))
        self.assertEqual(empty['type'], 'text')
        self.assertEqual([w['question_number'] for w in result['warnings']], [7])

    def test_invalid_structure(self):
        with self.assertRaises(ValueError):
            self.parse(b'{"no": "questions"}')
        with self.assertRaises(ValueError):
            self.parse(b'not json at all')


class TestCreatorImportApiTests(QuizTestBase):
    def upload(self, name, content, **extra):
        self.auth(self.teacher_token)
        return self.client.post('/api/v1/quizzes/import/', {
            'file': SimpleUploadedFile(name, content), **extra}, format='multipart')

    def test_json_import_creates_a_draft_with_points_and_multiple_answers(self):
        resp = self.upload('test.json', build_json('teacher'), topic='IELTS', course=self.course_id)
        self.assertEqual(resp.status_code, 201, resp.content)
        body = resp.json()
        self.assertEqual(body['status'], 'draft')
        self.assertEqual(body['warnings'], [])
        quiz = Quiz.objects.get(pk=body['id'])
        by_text = {q.text: q for q in quiz.questions.all()}
        self.assertEqual(by_text['Which are prime numbers?'].type, Question.Type.MULTIPLE)
        self.assertEqual(by_text['Which are prime numbers?'].points, 3)
        self.assertEqual(by_text['2 + 2 = ?'].points, 2)
        self.assertEqual(by_text['Which are prime numbers?'].options.filter(is_correct=True).count(), 2)
        # to'liq javobli qoralamani e'lon qilish mumkin
        self.auth(self.teacher_token)
        self.assertEqual(self.client.post(f'/api/v1/quizzes/{quiz.id}/publish/').status_code, 200)

    def test_json_preview_without_topic_saves_nothing(self):
        resp = self.upload('test.json', build_json('teacher'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.json()['questions']), 3)
        self.assertEqual(Quiz.objects.count(), 0)

    def test_docx_and_xlsx_exports_import_through_the_api(self):
        for name, buf in (('t.docx', build_docx('teacher')), ('t.xlsx', build_xlsx('teacher'))):
            resp = self.upload(name, buf.read())
            self.assertEqual(resp.status_code, 200, (name, resp.content))
            self.assertEqual(len(resp.json()['questions']), 3, name)
            self.assertEqual(resp.json()['warnings'], [], name)

    def test_broken_or_unknown_files_give_a_clear_error(self):
        self.assertEqual(self.upload('t.json', b'{"x": 1}').status_code, 400)
        self.assertEqual(self.upload('t.json', b'junk').status_code, 400)
        resp = self.upload('t.pdf', b'%PDF')
        self.assertEqual(resp.status_code, 400)
        self.assertIn('.json', str(resp.json()))


class NewExportFormatTests(SimpleTestCase):
    """Test-creator'ning yangi (2026-10) JSON eksporti: `text`, `number`, `options[].letter/is_correct`, `answer`."""

    PAYLOAD = {'title': 'IELTS Reading', 'mode': 'teacher', 'questions': [
        {'number': 2, 'text': '2 + 2 = ?', 'difficulty': 'easy', 'answer': 'B', 'explanation': '',
         'options': [{'letter': 'A', 'text': '3', 'is_correct': False}, {'letter': 'B', 'text': '4', 'is_correct': True}]},
        {'number': 1, 'text': 'Capital of France?', 'difficulty': 'easy', 'answer': 'A', 'explanation': 'Paris',
         'options': [{'letter': 'A', 'text': 'Paris', 'is_correct': True}, {'letter': 'B', 'text': 'Rome', 'is_correct': False},
                     {'letter': 'C', 'text': 'Berlin', 'is_correct': False}]},
    ]}

    def parse(self, payload=None):
        return test_creator_import.parse_test_creator_json(io.BytesIO(json.dumps(payload or self.PAYLOAD).encode()))

    def test_questions_are_read_in_order_with_the_right_answers(self):
        preview = self.parse()
        self.assertEqual([q['text'] for q in preview['questions']], ['Capital of France?', '2 + 2 = ?'])
        self.assertEqual([[o['is_correct'] for o in q['options']] for q in preview['questions']],
                         [[True, False, False], [False, True]])
        self.assertEqual((preview['title'], preview['warnings']), ('IELTS Reading', []))

    def test_a_letter_only_answer_key_is_enough(self):
        payload = json.loads(json.dumps(self.PAYLOAD))
        for question in payload['questions']:
            for option in question['options']:
                option.pop('is_correct')
        preview = self.parse(payload)
        self.assertEqual([[o['is_correct'] for o in q['options']] for q in preview['questions']],
                         [[True, False, False], [False, True]])

    def test_student_copy_without_answers_is_flagged(self):
        payload = json.loads(json.dumps(self.PAYLOAD))
        for question in payload['questions']:
            question.pop('answer')
            question['options'] = [{'letter': o['letter'], 'text': o['text']} for o in question['options']]
        preview = self.parse(payload)
        self.assertEqual({w['reason'] for w in preview['warnings']}, {'answer_not_detected'})
