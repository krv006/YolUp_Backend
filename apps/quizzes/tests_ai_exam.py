"""Imtihon qoidalari bo'yicha AI test generatori: matn olish, AI mijozi, reja, savol turlari, butun ish oqimi.
OpenAI chaqiruvlari soxta (`requests.post` almashtiriladi)."""
import io
import json
import shutil
import tempfile
import zipfile
from unittest.mock import MagicMock, patch

import requests
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import SimpleTestCase, override_settings

from apps.notifications.models import Notification

from . import ai_exam, llm, material_text
from .models import AiQuizJob, Quiz
from .tests_question_types import QuizTestBase

LLM_ON = dict(OPENAI_API_KEY='sk-test', OPENAI_MODEL='gpt-test', OPENAI_BASE_URL='https://llm.test/v1')
LONG = ('Honeybees are social insects that live in large colonies with one queen, thousands of workers and '
        'a few drones. ') * 3


def reply(status=200, body=None):
    mock = MagicMock()
    mock.status_code = status
    mock.json.return_value = body if body is not None else {}
    return mock


def completion(data, finish='stop'):
    text = data if isinstance(data, str) else json.dumps(data)
    return reply(200, {'choices': [{'message': {'content': text}, 'finish_reason': finish}],
                       'usage': {'total_tokens': 123}})


def error(status, code='', message='x'):
    return reply(status, {'error': {'code': code, 'message': message}})


# ─── Matn olish ────────────────────────────────────────────────────────────


def make_pdf(text: str) -> bytes:
    stream = f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'
    objects = [
        '<< /Type /Catalog /Pages 2 0 R >>',
        '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R '
        '/Resources << /Font << /F1 5 0 R >> >> >>',
        f'<< /Length {len(stream)} >>\nstream\n{stream}\nendstream',
        '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    out, offsets = '%PDF-1.4\n', []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f'{number} 0 obj\n{body}\nendobj\n'
    xref = len(out)
    out += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'
    out += ''.join(f'{offset:010d} 00000 n \n' for offset in offsets)
    out += f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'
    return out.encode('latin-1')


class MaterialTextTests(SimpleTestCase):
    def test_plain_text_and_markdown_and_csv(self):
        self.assertIn('Honeybees', material_text.extract_text(LONG.encode(), 'a.txt'))
        self.assertIn('Honeybees', material_text.extract_text(LONG.encode(), 'a.md'))
        csv_text = material_text.extract_text(('name,about\n' + 'bee,' + LONG + '\n').encode(), 'a.csv')
        self.assertIn('bee | Honeybees', csv_text)

    def test_pdf(self):
        self.assertIn('Honeybees', material_text.extract_text(make_pdf(LONG.replace('(', '').replace(')', '')), 'a.pdf'))

    def test_docx_includes_paragraphs_and_tables(self):
        from docx import Document

        document = Document()
        document.add_paragraph(LONG)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Cell A', 'Cell B'
        buffer = io.BytesIO()
        document.save(buffer)
        text = material_text.extract_text(buffer.getvalue(), 'a.docx')
        self.assertIn('Honeybees', text)
        self.assertIn('Cell A | Cell B', text)

    def test_pptx_slides_are_read_in_order(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            for number in (10, 2, 1):  # tartibsiz yozamiz — sonli tartibda o'qilishi kerak
                archive.writestr(
                    f'ppt/slides/slide{number}.xml',
                    '<p:sld xmlns:p="p" xmlns:a="a"><a:t>Slide %d %s</a:t></p:sld>' % (number, LONG))
        text = material_text.extract_text(buffer.getvalue(), 'a.pptx')
        self.assertLess(text.index('Slide 1 '), text.index('Slide 2 '))
        self.assertLess(text.index('Slide 2 '), text.index('Slide 10 '))

    def test_xlsx(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(['term', LONG])
        buffer = io.BytesIO()
        workbook.save(buffer)
        self.assertIn('term | Honeybees', material_text.extract_text(buffer.getvalue(), 'a.xlsx'))

    def test_scanned_or_empty_files_are_rejected_with_a_clear_message(self):
        with self.assertRaises(material_text.MaterialError) as ctx:
            material_text.extract_text(make_pdf('tiny'), 'scan.pdf')
        self.assertTrue(ctx.exception.permanent)
        self.assertIn('matn deyarli yo\'q', str(ctx.exception))

    def test_unsupported_and_broken_files(self):
        with self.assertRaises(material_text.MaterialError):
            material_text.extract_text(b'MZ', 'virus.exe')
        with self.assertRaises(material_text.MaterialError):
            material_text.extract_text(b'not a real docx', 'broken.docx')

    def test_text_is_truncated_to_the_limit(self):
        self.assertEqual(len(material_text.extract_text((LONG * 50).encode(), 'a.txt', max_chars=500)), 500)


# ─── AI mijozi ─────────────────────────────────────────────────────────────


@override_settings(**LLM_ON)
class LlmClientTests(SimpleTestCase):
    def call(self, *responses):
        with patch('apps.quizzes.llm.requests.post', side_effect=list(responses)) as post:
            return llm.chat_json('system JSON', 'user'), post

    def test_success_sends_model_json_mode_and_auth(self):
        data, post = self.call(completion({'ok': 1}))
        self.assertEqual(data, {'ok': 1})
        args, kwargs = post.call_args
        self.assertEqual(args[0], 'https://llm.test/v1/chat/completions')
        self.assertEqual(kwargs['json']['model'], 'gpt-test')
        self.assertEqual(kwargs['json']['response_format'], {'type': 'json_object'})
        self.assertEqual(kwargs['headers']['Authorization'], 'Bearer sk-test')
        self.assertEqual([m['role'] for m in kwargs['json']['messages']], ['system', 'user'])

    def test_markdown_fences_are_stripped(self):
        data, _ = self.call(completion('```json\n{"a": 2}\n```'))
        self.assertEqual(data, {'a': 2})

    def test_invalid_json_is_retried_once_with_a_correction(self):
        data, post = self.call(completion('nope'), completion({'a': 3}))
        self.assertEqual(data, {'a': 3})
        self.assertIn('not a valid', post.call_args.kwargs['json']['messages'][-1]['content'])

    def test_still_invalid_json_is_a_transient_error(self):
        with patch('apps.quizzes.llm.requests.post', side_effect=[completion('x'), completion('y')]):
            with self.assertRaises(llm.LLMError) as ctx:
                llm.chat_json('s', 'u')
        self.assertFalse(ctx.exception.permanent)

    def test_truncated_reply_is_reported_as_such(self):
        with patch('apps.quizzes.llm.requests.post',
                   side_effect=[completion('{"a": [1,', 'length'), completion('{"a": [', 'length')]):
            with self.assertRaises(llm.LLMError) as ctx:
                llm.chat_json('s', 'u')
        self.assertIn('qirqilib', str(ctx.exception))

    def test_permanent_errors(self):
        for response, fragment in (
            (error(401), 'kalit'), (error(403), 'kalit'),
            (error(429, 'insufficient_quota'), 'mablag'), (error(404), 'modeli topilmadi'),
            (error(400, message='bad param'), 'bad param'),
        ):
            with patch('apps.quizzes.llm.requests.post', return_value=response):
                with self.assertRaises(llm.LLMError) as ctx:
                    llm.chat_json('s', 'u')
            self.assertTrue(ctx.exception.permanent, fragment)
            self.assertIn(fragment, str(ctx.exception))

    def test_transient_errors(self):
        for response in (error(429, 'rate_limit_exceeded'), error(500), error(503)):
            with patch('apps.quizzes.llm.requests.post', return_value=response):
                with self.assertRaises(llm.LLMError) as ctx:
                    llm.chat_json('s', 'u')
            self.assertFalse(ctx.exception.permanent)
        with patch('apps.quizzes.llm.requests.post', side_effect=requests.ConnectionError('down')):
            with self.assertRaises(llm.LLMError) as ctx:
                llm.chat_json('s', 'u')
        self.assertFalse(ctx.exception.permanent)

    @override_settings(OPENAI_API_KEY='')
    def test_disabled_without_a_key(self):
        self.assertFalse(llm.enabled())


# ─── Reja va savol turlari ─────────────────────────────────────────────────


@override_settings(AI_EXAM_MAX_QUESTIONS=80)
class PlanTests(SimpleTestCase):
    def test_valid_plan_is_kept_and_clamped(self):
        plan = ai_exam.normalize_plan({
            'exam_name': 'IELTS Reading', 'language': 'en', 'rules': ['Answers follow the text order', ''],
            'sections': [{
                'title': 'Passage 1', 'passage': {'min_words': 700, 'max_words': 900},
                'blocks': [
                    {'type': 'true_false_not_given', 'count': 6, 'instruction': 'Do the statements agree?'},
                    {'type': 'fill_blank', 'count': 7, 'options': 99, 'points': 50},
                    {'type': 'poetry', 'count': 3},          # noma'lum tur — tashlanadi
                    {'type': 'mcq', 'count': 0},             # son yaroqsiz — tashlanadi
                ],
            }],
        }, fallback_count=20)
        section = plan['sections'][0]
        self.assertEqual([b['type'] for b in section['blocks']], ['true_false_not_given', 'fill_blank'])
        self.assertEqual((section['blocks'][1]['options'], section['blocks'][1]['points']), (6, 10))
        self.assertEqual(section['passage'], {'min_words': 700, 'max_words': 900, 'description': ''})
        self.assertEqual(plan['rules'], ['Answers follow the text order'])
        self.assertEqual(plan['language'], 'en')

    def test_total_questions_are_capped(self):
        plan = ai_exam.normalize_plan({'sections': [
            {'title': f'S{i}', 'blocks': [{'type': 'mcq', 'count': 40}]} for i in range(5)]}, fallback_count=20)
        self.assertEqual(sum(b['count'] for s in plan['sections'] for b in s['blocks']), 80)

    def test_garbage_falls_back_to_a_simple_plan(self):
        for raw in (None, 'text', {}, {'sections': 'x'}, {'sections': [{'blocks': [{'type': 'nope', 'count': 2}]}]}):
            plan = ai_exam.normalize_plan(raw, fallback_count=12)
            self.assertEqual(plan['sections'][0]['blocks'], [
                {'type': 'mcq', 'count': 12, 'instruction': '', 'options': 4, 'points': 1}])

    def test_without_rules_no_ai_call_is_made(self):
        with patch('apps.quizzes.llm.chat_json') as chat:
            plan = ai_exam.plan_exam('   ', target_total=15)
        chat.assert_not_called()
        self.assertEqual(plan['sections'][0]['blocks'][0]['count'], 15)


class QuestionConversionTests(SimpleTestCase):
    def convert(self, raw):
        question = ai_exam._to_question(raw, points=1)
        return ai_exam._validated(question) if question else None

    def test_mcq_and_multiple(self):
        single = self.convert({'kind': 'mcq', 'text': 'Q?', 'options': ['a', 'b', 'c'], 'answer': 1})
        self.assertEqual(single['type'], 'single')
        self.assertEqual([o['is_correct'] for o in single['options']], [False, True, False])
        multiple = self.convert({'kind': 'mcq_multiple', 'text': 'Pick two', 'options': ['a', 'b', 'c'],
                                 'answers': [0, 2, 9]})
        self.assertEqual([o['is_correct'] for o in multiple['options']], [True, False, True])

    def test_true_false_not_given_and_yes_no(self):
        tfng = self.convert({'kind': 'true_false_not_given', 'text': 'S', 'answer': 'Not Given'})
        self.assertEqual([o['text'] for o in tfng['options']], ['True', 'False', 'Not Given'])
        self.assertEqual([o['is_correct'] for o in tfng['options']], [False, False, True])
        yn = self.convert({'kind': 'yes_no_not_given', 'text': 'S', 'answer': 'yes'})
        self.assertEqual([o['text'] for o in yn['options']], ['Yes', 'No', 'Not Given'])
        self.assertTrue(yn['options'][0]['is_correct'])
        self.assertIsNone(self.convert({'kind': 'true_false_not_given', 'text': 'S', 'answer': 'maybe'}))

    def test_true_false(self):
        self.assertIs(self.convert({'kind': 'true_false', 'text': 'S', 'answer': 'false'})['correct_bool'], False)
        self.assertIsNone(self.convert({'kind': 'true_false', 'text': 'S', 'answer': 'maybe'}))

    def test_matching(self):
        question = self.convert({'kind': 'matching', 'text': 'Match', 'pairs': [
            {'left': 'Para A', 'right': 'Heading 1'}, {'left': 'Para B', 'right': 'Heading 2'}]})
        self.assertEqual(question['type'], 'matching')
        self.assertIsNone(self.convert({'kind': 'matching', 'text': 'Match', 'pairs': [{'left': 'A', 'right': 'B'}]}))
        self.assertIsNone(self.convert({'kind': 'matching', 'text': 'Match', 'pairs': [
            {'left': 'A', 'right': 'Same'}, {'left': 'B', 'right': 'Same'}]}))  # o'ng qiymat takrori — serializer rad etadi

    def test_fill_blank_needs_markers_in_order_and_answers(self):
        good = self.convert({'kind': 'fill_blank', 'text': 'The {{1}} lays {{2}}.',
                             'blanks': [['queen', 'Queen'], ['eggs']]})
        self.assertEqual([b['answers'] for b in good['blanks']], [['queen', 'Queen'], ['eggs']])
        self.assertIsNone(self.convert({'kind': 'fill_blank', 'text': 'No gaps here', 'blanks': [['x']]}))
        self.assertIsNone(self.convert({'kind': 'fill_blank', 'text': '{{2}} only', 'blanks': [['x'], ['y']]}))
        self.assertIsNone(self.convert({'kind': 'fill_blank', 'text': '{{1}} {{2}}', 'blanks': [['x']]}))

    def test_short_answer_numeric_ordering(self):
        self.assertEqual(self.convert({'kind': 'short_answer', 'text': 'Q', 'answers': ['bees']})['accepted_answers'],
                         ['bees'])
        self.assertEqual(self.convert({'kind': 'numeric', 'text': 'Q', 'answers': ['12'], 'tolerance': 0.5})['tolerance'], 0.5)
        self.assertIsNone(self.convert({'kind': 'numeric', 'text': 'Q', 'answers': ['twelve']}))
        self.assertEqual(self.convert({'kind': 'ordering', 'text': 'Q', 'items': ['a', 'b', 'c']})['type'], 'ordering')
        self.assertIsNone(self.convert({'kind': 'ordering', 'text': 'Q', 'items': ['only one']}))

    def test_garbage_is_rejected(self):
        for raw in (None, 'x', {}, {'kind': 'mcq'}, {'kind': 'poem', 'text': 'Q'}, {'kind': 'mcq', 'text': ''},
                    {'kind': 'mcq', 'text': 'Q', 'options': ['a'], 'answer': 0},
                    {'kind': 'mcq', 'text': 'Q', 'options': ['a', 'b'], 'answer': 5}):
            self.assertIsNone(self.convert(raw), raw)


# ─── Butun ish oqimi ───────────────────────────────────────────────────────

IELTS_RULES = (
    'IELTS Academic Reading: 3 passages of 700-900 words. Passage 1: questions 1-4 True/False/Not Given, '
    'questions 5-7 sentence completion (no more than two words). Passage 2: questions 8-9 multiple choice.'
)


def planner_reply(**extra):
    return {
        'exam_name': 'IELTS Reading', 'language': 'en', 'rules': ['Answers follow the order of the text'],
        'sections': [
            {'title': 'Passage 1', 'passage': {'min_words': 700, 'max_words': 900}, 'blocks': [
                {'type': 'true_false_not_given', 'count': 2, 'instruction': 'Do the statements agree with the text?'},
                {'type': 'fill_blank', 'count': 1, 'instruction': 'Complete the sentence. NO MORE THAN TWO WORDS.'},
            ]},
            {'title': 'Passage 2', 'passage': {'min_words': 700, 'max_words': 900}, 'blocks': [
                {'type': 'mcq', 'count': 2, 'instruction': 'Choose the correct letter.'},
            ]},
        ], **extra,
    }


def section_reply(number, *, good=True):
    if number == 1:
        questions = [
            {'kind': 'true_false_not_given', 'text': 'Bees live alone.', 'answer': 'false'},
            {'kind': 'true_false_not_given', 'text': 'A colony has one queen.', 'answer': 'true'},
            {'kind': 'fill_blank', 'text': 'Workers number {{1}}.', 'blanks': [['thousands']]},
        ]
    else:
        questions = [
            {'kind': 'mcq', 'text': 'How many queens?', 'options': ['one', 'two', 'ten'], 'answer': 0},
            {'kind': 'mcq', 'text': 'Broken question', 'options': ['only one'], 'answer': 0},  # yaroqsiz — tashlanadi
        ]
    return {'passage': LONG * 2 if good else '', 'questions': questions}


class FakeOpenAI:
    """Planner yoki writer chaqiruviga qarab javob beradi; chaqiruvlarni yozib boradi."""

    def __init__(self):
        self.calls = []
        self.fail_section = None  # shu bo'lim raqami uchun vaqtincha xato

    def __call__(self, url, json=None, **kwargs):
        system, user = json['messages'][0]['content'], json['messages'][1]['content']
        if system.startswith('You are an exam designer'):
            self.calls.append(('plan', None))
            return completion(planner_reply())
        spec = __import__('json').loads(user)['section_spec']
        number = spec['section_number']
        self.calls.append(('section', number))
        if self.fail_section == number:
            return error(503)
        return completion(section_reply(number))


@override_settings(**LLM_ON)
class AiExamJobTests(QuizTestBase):
    def setUp(self):
        super().setUp()
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)
        self.fake = FakeOpenAI()
        patcher = patch('apps.quizzes.llm.requests.post', side_effect=self.fake)
        patcher.start()
        self.addCleanup(patcher.stop)

    def upload(self, name='material.txt', text=LONG * 5):
        return SimpleUploadedFile(name, text.encode(), content_type='text/plain')

    def start(self, **extra):
        self.auth(self.teacher_token)
        data = {'file': self.upload(), 'topic': 'Bees', 'question_count': 10, 'course': self.course_id, **extra}
        return self.client.post('/api/v1/quizzes/ai-generate/', data, format='multipart')

    def sync(self):
        call_command('sync_ai_quizzes', stdout=io.StringIO())

    def test_rules_file_produces_a_sectioned_mixed_type_draft(self):
        job_id = self.start(rules_file=self.upload('ielts_rules.txt', IELTS_RULES * 3)).json()['id']
        self.assertEqual(AiQuizJob.objects.get(pk=job_id).standard, '')
        self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual(job.status, 'done', job.error)
        quiz = Quiz.objects.get(pk=job.quiz_id)
        self.assertEqual(quiz.status, Quiz.Status.DRAFT)
        self.assertEqual(quiz.groups.count(), 2)  # har matn parchasi — guruh
        questions = list(quiz.questions.order_by('order'))
        self.assertEqual([q.type for q in questions], ['single', 'single', 'fill_blank', 'single'])  # tashlangani yo'q
        self.assertEqual([q.group.title for q in questions], ['Passage 1'] * 3 + ['Passage 2'])
        # ko'rsatma har blokning birinchi savoliga qo'shilgan, tartib saqlangan
        self.assertTrue(questions[0].text.startswith('Do the statements agree with the text?\n\nBees live alone.'))
        self.assertEqual(questions[1].text, 'A colony has one queen.')
        self.assertTrue(questions[2].text.startswith('Complete the sentence. NO MORE THAN TWO WORDS.'))
        self.assertEqual([o.text for o in questions[0].options.order_by('order')], ['True', 'False', 'Not Given'])
        self.assertTrue(questions[0].options.get(text='False').is_correct)
        self.assertEqual(job.summary.split(' (')[0], "2 bo'lim · 4 savol")
        self.assertIn('tushib qoldi', job.summary)
        # 1 reja + 2 bo'lim, boshqa hech narsa; fayllar o'chirilgan; bildirishnoma bor
        self.assertEqual([c[0] for c in self.fake.calls], ['plan', 'section', 'section'])
        self.assertFalse(job.source_file or job.rules_file)
        self.assertEqual(Notification.objects.get(kind='ai_quiz_ready').link_id, str(quiz.id))

    def test_rules_as_pasted_text_work_the_same(self):
        job_id = self.start(rules_text=IELTS_RULES).json()['id']
        self.sync()
        self.assertEqual(AiQuizJob.objects.get(pk=job_id).status, 'done')
        plan_call = [c for c in self.fake.calls if c[0] == 'plan']
        self.assertEqual(len(plan_call), 1)

    def test_without_rules_a_simple_multiple_choice_test_is_made_with_one_ai_call(self):
        with patch('apps.quizzes.llm.requests.post') as post:
            post.return_value = completion({'passage': '', 'questions': [
                {'kind': 'mcq', 'text': f'Q{i}?', 'options': ['a', 'b', 'c', 'd'], 'answer': i % 4}
                for i in range(10)]})
            job_id = self.start().json()['id']
            self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual(job.status, 'done', job.error)
        self.assertEqual(post.call_count, 1)
        quiz = Quiz.objects.get(pk=job.quiz_id)
        self.assertEqual(quiz.questions.count(), 10)
        self.assertEqual(quiz.groups.count(), 0)
        self.assertEqual(job.summary, "1 bo'lim · 10 savol")

    def test_serializer_reports_the_mode(self):
        simple = self.start().json()
        rules = self.start(rules_text=IELTS_RULES).json()
        self.assertEqual((simple['mode'], rules['mode']), ('simple', 'rules'))

    def test_a_failed_section_resumes_without_redoing_finished_ones(self):
        self.fake.fail_section = 2
        job_id = self.start(rules_text=IELTS_RULES).json()['id']
        self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual((job.status, job.attempts), ('generating', 1))
        self.assertIn('0', job.plan['sections'])          # 1-bo'lim saqlangan
        self.assertNotIn('1', job.plan['sections'])
        self.fake.fail_section = None
        AiQuizJob.objects.filter(pk=job_id).update(updated_at=job.updated_at.replace(year=2020))
        self.sync()
        self.assertEqual(AiQuizJob.objects.get(pk=job_id).status, 'done')
        self.assertEqual([c for c in self.fake.calls if c[0] == 'plan'], [('plan', None)])
        self.assertEqual([c[1] for c in self.fake.calls if c[0] == 'section'], [1, 2, 2])  # 1-bo'lim qayta yozilmagan

    def test_missing_passage_is_retried_as_a_transient_error(self):
        with patch('apps.quizzes.ai_exam._WRITER_SYSTEM', ai_exam._WRITER_SYSTEM):
            self.fake_bad = patch('apps.quizzes.llm.requests.post', side_effect=lambda *a, **k: (
                completion(planner_reply()) if k['json']['messages'][0]['content'].startswith('You are an exam designer')
                else completion(section_reply(1, good=False))))
            with self.fake_bad:
                job_id = self.start(rules_text=IELTS_RULES).json()['id']
                self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual((job.status, job.attempts), ('generating', 1))

    def test_invalid_key_fails_immediately_with_a_clear_message(self):
        with patch('apps.quizzes.llm.requests.post', return_value=error(401)):
            job_id = self.start(rules_text=IELTS_RULES).json()['id']
            self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual(job.status, 'failed')
        self.assertIn('OPENAI_API_KEY', job.error)
        self.assertTrue(Notification.objects.filter(kind='ai_quiz_failed').exists())

    def post_raw(self, **fields):
        self.auth(self.teacher_token)
        data = {'topic': 'Bees', 'question_count': 10, 'course': self.course_id, **fields}
        return self.client.post('/api/v1/quizzes/ai-generate/', data, format='multipart')

    def test_scanned_material_fails_in_the_background_with_a_helpful_message(self):
        scanned = SimpleUploadedFile('scan.pdf', make_pdf('tiny'), content_type='application/pdf')
        resp = self.post_raw(file=scanned)
        self.assertEqual(resp.status_code, 202)  # so'rov faylni o'qimaydi — darhol qaytadi
        job_id = resp.json()['id']
        stored = AiQuizJob.objects.get(pk=job_id).source_paths
        self.assertEqual(len(stored), 1)
        self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual(job.status, 'failed')
        self.assertIn("matn deyarli yo'q", job.error)
        self.assertFalse(default_storage.exists(stored[0]['path']))  # fayl tozalangan
        self.assertEqual(job.source_paths, [])

    def test_request_returns_without_reading_the_files(self):
        with patch('apps.quizzes.material_text.extract_text') as extract:
            resp = self.post_raw(file=self.upload())
        self.assertEqual(resp.status_code, 202)
        extract.assert_not_called()

    def test_several_files_are_combined_into_one_material(self):
        files = [self.upload(f'part{i}.txt', f'Section {i} about bees. ' * 20) for i in (1, 2, 3)]
        job_id = self.post_raw(file=files, rules_text=IELTS_RULES).json()['id']
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual(job.source_name, 'part1.txt, part2.txt, part3.txt')
        stored = list(job.source_paths)
        self.assertEqual([e['name'] for e in stored], ['part1.txt', 'part2.txt', 'part3.txt'])
        self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual(job.status, 'done')
        for i in (1, 2, 3):
            self.assertIn(f'=== part{i}.txt ===', job.plan['material'])
        self.assertTrue(all(not default_storage.exists(e['path']) for e in stored))

    def test_pasted_text_works_without_any_file(self):
        job_id = self.post_raw(material_text='Bees are fascinating insects. ' * 20).json()['id']
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertIn('Bees are fascinating', job.source_text)
        self.assertEqual(job.source_name, '')

    def test_material_is_optional_ai_writes_from_the_topic_and_rules(self):
        seen = {}

        def capture(url, json=None, **kwargs):
            seen.setdefault('calls', []).append(json['messages'][1]['content'])
            return self.fake(url, json=json, **kwargs)

        with patch('apps.quizzes.llm.requests.post', side_effect=capture):
            job_id = self.post_raw(rules_text=IELTS_RULES, topic='Climate change').json()['id']
            self.sync()
        self.assertEqual(AiQuizJob.objects.get(pk=job_id).status, 'done')
        section_prompts = [json.loads(c) for c in seen['calls'] if 'section_spec' in c]
        self.assertTrue(section_prompts)
        for payload in section_prompts:
            self.assertEqual((payload['source_material'], payload['topic']), ('', 'Climate change'))

    def test_too_many_files_and_one_bad_file(self):
        many = [self.upload(f'f{i}.txt') for i in range(6)]
        self.assertEqual(self.post_raw(file=many).status_code, 400)
        bad = SimpleUploadedFile('virus.exe', b'MZ')
        self.assertEqual(self.post_raw(file=[self.upload(), bad]).status_code, 400)

    def test_old_test_creator_path_still_needs_exactly_one_file(self):
        with override_settings(TEST_CREATOR_URL='http://tc-api:8000/api/v1', TEST_CREATOR_EMAIL='a@b.uz',
                               TEST_CREATOR_PASSWORD='x' * 12):
            self.assertEqual(self.post_raw(standard='uzbmb').status_code, 400)
            self.assertEqual(self.post_raw(standard='uzbmb', file=[self.upload(), self.upload('b.txt')]).status_code, 400)

    def test_all_questions_invalid_fails(self):
        with patch('apps.quizzes.llm.requests.post', return_value=completion({'passage': '', 'questions': [
                {'kind': 'mcq', 'text': 'Q', 'options': ['only'], 'answer': 0}]})):
            job_id = self.start().json()['id']
            self.sync()
        job = AiQuizJob.objects.get(pk=job_id)
        self.assertEqual(job.status, 'failed')
        self.assertIn('yaroqli savol', job.error)

    def test_rules_file_validation(self):
        self.assertEqual(self.start(rules_file=self.upload('rules.exe', IELTS_RULES)).status_code, 400)
        big = SimpleUploadedFile('rules.txt', b'x' * (6 * 1024 * 1024))
        self.assertEqual(self.start(rules_file=big).status_code, 400)
        self.assertEqual(self.start(rules_text='x' * 30001).status_code, 400)

    @override_settings(OPENAI_API_KEY='')
    def test_disabled_without_an_ai_key(self):
        resp = self.start()
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(AiQuizJob.objects.count(), 0)

    def test_old_test_creator_path_still_needs_its_own_service(self):
        resp = self.start(standard='uzbmb')  # TEST_CREATOR_URL sozlanmagan
        self.assertEqual(resp.status_code, 400)

    def test_engine_jobs_are_not_touched_when_ai_is_off(self):
        job_id = self.start(rules_text=IELTS_RULES).json()['id']
        with override_settings(OPENAI_API_KEY=''):
            self.sync()
        self.assertEqual(AiQuizJob.objects.get(pk=job_id).status, 'queued')

    # ── imtihon nomi bo'yicha (qoidalarsiz) ───────────────────────────

    def test_exam_name_alone_lets_the_ai_recall_the_official_format(self):
        seen = {}

        def capture(url, json=None, **kwargs):
            seen.setdefault('user', []).append(json['messages'][1]['content'])
            return self.fake(url, json=json, **kwargs)

        with patch('apps.quizzes.llm.requests.post', side_effect=capture):
            response = self.post_raw(exam_name='IELTS Academic Reading')
            self.assertEqual(response.json()['mode'], 'rules')
            self.sync()
        job = AiQuizJob.objects.get(pk=response.json()['id'])
        self.assertEqual((job.status, job.exam_name), ('done', 'IELTS Academic Reading'))
        planner_prompt = seen['user'][0]
        self.assertIn('exam: IELTS Academic Reading', planner_prompt)
        self.assertIn('none provided', planner_prompt)
        self.assertEqual([c[0] for c in self.fake.calls], ['plan', 'section', 'section'])
        self.assertEqual(Quiz.objects.get(pk=job.quiz_id).groups.count(), 2)

    def test_rules_override_the_exam_name(self):
        seen = {}

        def capture(url, json=None, **kwargs):
            seen.setdefault('user', []).append(json['messages'][1]['content'])
            return self.fake(url, json=json, **kwargs)

        with patch('apps.quizzes.llm.requests.post', side_effect=capture):
            self.post_raw(exam_name='IELTS', rules_text=IELTS_RULES)
            self.sync()
        self.assertIn('override your own knowledge', seen['user'][0])
        self.assertIn(IELTS_RULES[:40], seen['user'][0])

    def test_neither_name_nor_rules_means_no_planning_call(self):
        self.post_raw(topic='Bees')
        with patch('apps.quizzes.llm.chat_json') as chat:
            chat.return_value = {'passage': '', 'questions': []}
            self.sync()
        planner_calls = [c for c in chat.call_args_list if c.args and 'exam designer' in c.args[0]]
        self.assertEqual(planner_calls, [])
