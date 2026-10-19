"""Test natijalari: urinish raqami, avatar, urinmaganlar ro'yxati va serverdan xlsx/csv eksport."""
import csv
import io

import openpyxl
from django.core.files.base import ContentFile

from apps.accounts.models import User
from apps.lessons.models import Enrollment
from apps.quizzes.models import Option

from .tests_question_types import QuizTestBase


class AttemptsReportTests(QuizTestBase):
    def setUp(self):
        super().setUp()
        questions = [
            {'type': 'single', 'text': 'Birinchi', 'points': 2, 'options': [
                {'text': "to'g'ri", 'is_correct': True}, {'text': 'xato', 'is_correct': False}]},
            {'type': 'numeric', 'text': 'Ikkinchi', 'points': 3, 'accepted_answers': ['5']},
        ]
        resp = self.make_quiz(questions)
        self.assertEqual(resp.status_code, 201, resp.content)
        self.quiz_id = resp.json()['id']
        self.single = next(q for q in resp.json()['questions'] if q['type'] == 'single')
        self.numeric = next(q for q in resp.json()['questions'] if q['type'] == 'numeric')
        self.right = str(Option.objects.get(question_id=self.single['id'], is_correct=True).id)
        self.wrong = str(Option.objects.get(question_id=self.single['id'], is_correct=False).id)
        User.objects.filter(username='s1').update(first_name='Sardor', last_name='Aliyev')

    def attempt(self, correct):
        answers = [
            {'question': self.single['id'], 'selected_option': self.right if correct else self.wrong},
            {'question': self.numeric['id'], 'value_text': '5' if correct else '9'},
        ]
        resp = self.submit(self.quiz_id, answers)
        self.assertEqual(resp.status_code, 201, resp.content)
        return resp.json()

    def test_attempts_list_has_numbers_names_and_avatar(self):
        self.attempt(False)
        self.attempt(True)
        self.auth(self.teacher_token)
        rows = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/').json()
        self.assertEqual(sorted(r['attempt_no'] for r in rows), [1, 2])
        newest = max(rows, key=lambda r: r['created_at'])
        self.assertEqual((newest['attempt_no'], newest['student_full_name'], newest['student_avatar_url']), (2, 'Sardor Aliyev', None))
        self.assertEqual(newest['student_name'], 's1')  # eski maydon o'zgarmagan
        student = User.objects.get(username='s1')
        student.avatar.save('a.png', ContentFile(b'x'), save=True)
        rows = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/').json()
        self.assertTrue(rows[0]['student_avatar_url'].startswith('http'))
        self.auth(self.child_token)  # o'quvchi faqat o'zini ko'radi, raqami baribir to'g'ri
        own = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/').json()
        self.assertEqual(sorted(r['attempt_no'] for r in own), [1, 2])

    def test_xlsx_export_has_every_attempt_and_answer(self):
        self.attempt(False)
        self.attempt(True)
        self.auth(self.teacher_token)
        resp = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/export/?type=xlsx')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertIn('.xlsx', resp['Content-Disposition'])
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content))['Natijalar']
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][:8], ('№', 'Ism', 'Login', 'Urinish', 'Sana', 'Ball', 'Maksimal ball', 'Foiz'))
        self.assertEqual(rows[0][8:], ('S1', 'S2'))
        self.assertEqual(len(rows), 3)
        first, second = rows[1], rows[2]
        self.assertEqual((first[1], first[2], first[3], first[5], first[6], first[7]), ('Sardor Aliyev', 's1', 1, 0, 5, 0))
        self.assertEqual((second[3], second[5], second[7]), (2, 5, 100))
        self.assertIn('(2/2)', second[8])
        self.assertIn('(0/2)', first[8])
        legend = openpyxl.load_workbook(io.BytesIO(resp.content))['Savollar']
        self.assertEqual([c.value for c in legend[2]], ['S1', 'Birinchi', 2])

    def test_csv_export_and_defaults(self):
        self.attempt(True)
        self.auth(self.teacher_token)
        resp = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/export/?type=csv')
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.startswith(b'\xef\xbb\xbf'))
        rows = list(csv.reader(io.StringIO(resp.content.decode('utf-8-sig'))))
        self.assertEqual(rows[0][:3], ['№', 'Ism', 'Login'])
        self.assertEqual(rows[1][1], 'Sardor Aliyev')
        default = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/export/')
        self.assertIn('.xlsx', default['Content-Disposition'])
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/export/?type=pdf').status_code, 400)

    def test_export_is_for_the_teacher_only(self):
        self.auth(self.child_token)
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/export/').status_code, 403)
        self.auth(self.parent_token)
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{self.quiz_id}/attempts/export/').status_code, 403)

    def test_not_attempted_lists_enrolled_students_without_an_attempt(self):
        self.auth(self.teacher_token)
        before = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/not-attempted/').json()
        self.assertEqual([r['username'] for r in before], ['s1'])
        self.attempt(True)
        self.auth(self.teacher_token)
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{self.quiz_id}/not-attempted/').json(), [])
        late = User.objects.create_user(username='late', password='x', role=User.Role.STUDENT)
        Enrollment.objects.create(course_id=self.course_id, student=late, status='approved')
        Enrollment.objects.create(
            course_id=self.course_id, student=User.objects.create_user(username='wait', password='x', role=User.Role.STUDENT),
            status='pending',
        )
        rows = self.client.get(f'/api/v1/quizzes/{self.quiz_id}/not-attempted/').json()
        self.assertEqual([r['username'] for r in rows], ['late'])
        self.auth(self.child_token)
        self.assertEqual(self.client.get(f'/api/v1/quizzes/{self.quiz_id}/not-attempted/').status_code, 403)
