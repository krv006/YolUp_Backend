"""Test natijalarini serverdan yuklab olish (xlsx / csv): har o'quvchining har urinishi va har savol bo'yicha javobi."""
import csv
import io

from django.utils import timezone

FORMATS = {
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'csv': 'text/csv; charset=utf-8',
}


def _fmt_points(value: float) -> str:
    return f'{value:g}'


def _attempt_rows(quiz, attempts) -> tuple:
    """(sarlavha, qatorlar, savollar): qatorlar — ism bo'yicha, har o'quvchining urinishlari xronologik."""
    questions = list(quiz.questions.order_by('order', 'created_at'))
    header = ['№', 'Ism', 'Login', 'Urinish', 'Sana', 'Ball', 'Maksimal ball', 'Foiz'] + [
        f'S{i}' for i in range(1, len(questions) + 1)
    ]
    ordered = sorted(
        attempts, key=lambda a: (a.student.get_full_name().lower() or a.student.username.lower(), a.created_at),
    )
    counters: dict = {}
    rows = []
    for number, attempt in enumerate(ordered, 1):
        counters[attempt.student_id] = counters.get(attempt.student_id, 0) + 1
        given = {answer.question_id: answer for answer in attempt.answers.all()}
        cells = []
        for question in questions:
            answer = given.get(question.id)
            if answer is None:
                cells.append('—')
            else:
                text = answer.given_display or '—'
                cells.append(f'{text} ({_fmt_points(answer.earned_points)}/{question.points})')
        percent = round(attempt.score * 100 / attempt.max_score) if attempt.max_score else 0
        rows.append([
            number, attempt.student.get_full_name() or attempt.student.username, attempt.student.username,
            counters[attempt.student_id], timezone.localtime(attempt.created_at).strftime('%Y-%m-%d %H:%M'),
            round(attempt.score, 2), attempt.max_score, percent, *cells,
        ])
    return header, rows, questions


def build_attempts_export(quiz, attempts, fmt: str) -> tuple:
    """Qaytaradi: (baytlar, content_type)."""
    if fmt not in FORMATS:
        raise ValueError(f"Format qo'llab-quvvatlanmaydi: {fmt}. Mumkin: {', '.join(FORMATS)}.")
    header, rows, questions = _attempt_rows(quiz, attempts)
    if fmt == 'csv':
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(header)
        writer.writerows(rows)
        return buffer.getvalue().encode('utf-8-sig'), FORMATS['csv']  # BOM — Excel o'zbekcha harflarni to'g'ri ochadi

    import openpyxl
    from openpyxl.styles import Alignment, Font

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Natijalar'
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column, width in zip('ABCDEFGH', (5, 28, 18, 9, 17, 8, 14, 8)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'C2'
    legend = workbook.create_sheet('Savollar')
    legend.append(['Savol', 'Matn', 'Ball'])
    for cell in legend[1]:
        cell.font = Font(bold=True)
    for index, question in enumerate(questions, 1):
        legend.append([f'S{index}', question.text, question.points])
    legend.column_dimensions['B'].width = 80
    for row in legend.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue(), FORMATS['xlsx']
