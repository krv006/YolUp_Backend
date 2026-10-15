"""Testni PDF / Word / Excel / CSV / JSON qilib yuklab olish.

Ikki rejim: `student` — o'quvchiga tarqatish uchun (javob kaliti YO'Q), `teacher` — o'qituvchi uchun (to'g'ri javoblar
belgilangan, Listening skripti bilan). Manba — `QuizDetailSerializer` ma'lumoti (javob kaliti bilan), shuning uchun savol
turlari bitta joyda (`_question_block`) tushuntiriladi, formatlar esa faqat chizadi.
"""
import csv
import io
import json
import re
from xml.sax.saxutils import escape

FORMATS = {
    'pdf': 'application/pdf',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'csv': 'text/csv; charset=utf-8',
    'json': 'application/json',
}
MODES = ('student', 'teacher')
LETTERS = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'
_GAP = re.compile(r'\{\{(\d+)\}\}')


def _letter(index: int) -> str:
    return LETTERS[index] if index < len(LETTERS) else str(index + 1)


def _shuffled(items: list) -> list:
    """Tartibni qat'iy (har safar bir xil) aralashtiradi — tarqatma varaqda javob ketma-ketligi ko'rinib qolmasin."""
    return sorted(items, key=lambda value: (hash_text(value), value))


def hash_text(value: str) -> int:
    total = 0
    for char in str(value):
        total = (total * 131 + ord(char)) % 1000003
    return total


def _answer_list(values) -> str:
    return '; '.join(str(v) for v in values if str(v).strip())


def _question_block(question: dict, number: int, *, teacher: bool) -> dict:
    """Bitta savol: `lines` — [(matn, to'g'ri_mi)], `answer` — o'qituvchi uchun javob matni (o'quvchi rejimida bo'sh)."""
    qtype, lines, answer = question['type'], [], ''
    text = question['text']
    if qtype in ('single', 'multiple'):
        for index, option in enumerate(question.get('options') or []):
            lines.append((f"{_letter(index)}) {option['text']}", bool(option.get('is_correct'))))
        answer = _answer_list(_letter(i) for i, o in enumerate(question.get('options') or []) if o.get('is_correct'))
        if qtype == 'multiple':
            text += '  (bir nechta to\'g\'ri javob)'
    elif qtype == 'true_false':
        lines = [("Rost", question.get('correct_bool') is True), ("Yolg'on", question.get('correct_bool') is False)]
        answer = 'Rost' if question.get('correct_bool') else "Yolg'on"
    elif qtype == 'matching':
        pairs = question.get('pairs') or []
        rights = _shuffled([p['right'] for p in pairs])
        for index, pair in enumerate(pairs, start=1):
            lines.append((f"{index}. {pair['left']}", False))
        for index, right in enumerate(rights):
            lines.append((f"   {_letter(index)}) {right}", False))
        answer = _answer_list(f"{i}-{_letter(rights.index(p['right']))}" for i, p in enumerate(pairs, start=1))
    elif qtype == 'ordering':
        items = question.get('items') or []
        for index, item in enumerate(_shuffled(items), start=1):
            lines.append((f"{index}. {item}", False))
        answer = ' → '.join(items)
        text += '  (to\'g\'ri tartibga qo\'ying)'
    elif qtype == 'fill_blank':
        text = _GAP.sub(lambda m: f'______({m.group(1)})', text)
        blanks = question.get('blanks') or []
        answer = _answer_list(f"{i}) {' / '.join(b.get('answers', []))}" for i, b in enumerate(blanks, start=1))
    else:  # text, numeric
        lines.append(('Javob: ____________________', False))
        answer = _answer_list(question.get('accepted_answers') or [])
    return {'n': number, 'text': text, 'points': question.get('points'), 'lines': lines,
            'answer': answer if teacher else '', 'type': qtype}


def build_outline(data: dict, *, mode: str) -> list:
    """Chizish uchun ketma-ket bloklar: {'kind': 'group'|'question', ...}."""
    teacher = mode == 'teacher'
    groups = {str(g['id']): g for g in data.get('groups') or []}
    blocks, shown, number = [], set(), 0
    for question in sorted(data.get('questions') or [], key=lambda q: q.get('order', 0)):
        group_id = str(question['group']) if question.get('group') else None
        if group_id and group_id in groups and group_id not in shown:
            shown.add(group_id)
            group = groups[group_id]
            blocks.append({
                'kind': 'group', 'title': group.get('title') or '', 'passage': group.get('passage') or '',
                'audio': bool(group.get('audio_url')), 'transcript': (group.get('transcript') or '') if teacher else '',
            })
        number += 1
        blocks.append({'kind': 'question', **_question_block(question, number, teacher=teacher)})
    return blocks


def _title(data: dict) -> str:
    return data.get('title') or data.get('topic') or 'Test'


def _meta(data: dict, mode: str) -> str:
    parts = [data.get('subject_label') or '', "O'qituvchi nusxasi (javoblar bilan)" if mode == 'teacher' else 'Savollar']
    return ' · '.join(p for p in parts if p)


# ─── Word ────────────────────────────────────────────────────────────────────


def to_docx(data: dict, mode: str) -> bytes:
    import docx
    from docx.shared import RGBColor

    document = docx.Document()
    document.add_heading(_title(data), level=1)
    document.add_paragraph(_meta(data, mode))
    for block in build_outline(data, mode=mode):
        if block['kind'] == 'group':
            if block['title']:
                document.add_heading(block['title'], level=2)
            if block['passage']:
                document.add_paragraph(block['passage'])
            if block['audio']:
                document.add_paragraph('🔊 Audio: platformada ijro etiladi.')
            continue
        points = f" ({block['points']} ball)" if block['points'] else ''
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{block['n']}. {block['text']}").bold = True
        paragraph.add_run(points)
        for text, correct in block['lines']:
            line = document.add_paragraph()
            run = line.add_run(f"{text}{'  ✓' if (correct and mode == 'teacher') else ''}")
            if correct and mode == 'teacher':
                run.bold = True
                run.font.color.rgb = RGBColor(0x1B, 0x7F, 0x3B)
        if block['answer']:
            answer = document.add_paragraph()
            run = answer.add_run(f"Javob: {block['answer']}")
            run.italic = True
            run.font.color.rgb = RGBColor(0x1B, 0x7F, 0x3B)
    if mode == 'teacher':
        for block in build_outline(data, mode=mode):
            if block['kind'] == 'group' and block['transcript']:
                document.add_heading(f"Audio skripti{': ' + block['title'] if block['title'] else ''}", level=2)
                document.add_paragraph(block['transcript'])
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ─── PDF ─────────────────────────────────────────────────────────────────────

_FONT = {}


def _pdf_fonts():
    """O'zbekcha belgilarni (oʻ, gʻ, ʼ) chizadigan Unicode shrift: matplotlib bilan keladigan DejaVu Sans."""
    if _FONT:
        return _FONT
    from pathlib import Path

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular, bold = 'Helvetica', 'Helvetica-Bold'
    try:
        import matplotlib

        base = Path(matplotlib.get_data_path()) / 'fonts' / 'ttf'
        if (base / 'DejaVuSans.ttf').exists():
            pdfmetrics.registerFont(TTFont('QuizSans', str(base / 'DejaVuSans.ttf')))
            regular = bold = 'QuizSans'
            if (base / 'DejaVuSans-Bold.ttf').exists():
                pdfmetrics.registerFont(TTFont('QuizSans-Bold', str(base / 'DejaVuSans-Bold.ttf')))
                bold = 'QuizSans-Bold'
    except Exception:  # noqa: BLE001 — shrift topilmasa ham PDF chiqadi (standart shrift bilan)
        pass
    _FONT.update(regular=regular, bold=bold)
    return _FONT


def to_pdf(data: dict, mode: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    fonts = _pdf_fonts()
    body = ParagraphStyle('body', fontName=fonts['regular'], fontSize=10.5, leading=14, spaceAfter=3)
    bold = ParagraphStyle('bold', parent=body, fontName=fonts['bold'], spaceBefore=8)
    h1 = ParagraphStyle('h1', parent=body, fontName=fonts['bold'], fontSize=17, leading=21, spaceAfter=4)
    h2 = ParagraphStyle('h2', parent=body, fontName=fonts['bold'], fontSize=13, leading=17, spaceBefore=10)
    meta = ParagraphStyle('meta', parent=body, textColor='#666666', spaceAfter=8)
    option = ParagraphStyle('option', parent=body, leftIndent=14)
    correct = ParagraphStyle('correct', parent=option, fontName=fonts['bold'], textColor='#1b7f3b')
    key = ParagraphStyle('key', parent=body, leftIndent=14, textColor='#1b7f3b')

    def para(text, style):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), style)

    story = [para(_title(data), h1), para(_meta(data, mode), meta)]
    for block in build_outline(data, mode=mode):
        if block['kind'] == 'group':
            if block['title']:
                story.append(para(block['title'], h2))
            if block['passage']:
                story.append(para(block['passage'], body))
            if block['audio']:
                story.append(para('Audio: platformada ijro etiladi.', meta))
            continue
        points = f" ({block['points']} ball)" if block['points'] else ''
        story.append(para(f"{block['n']}. {block['text']}{points}", bold))
        for text, is_correct in block['lines']:
            marked = is_correct and mode == 'teacher'
            story.append(para(f"{text}{'  ✓' if marked else ''}", correct if marked else option))
        if block['answer']:
            story.append(para(f"Javob: {block['answer']}", key))
    if mode == 'teacher':
        for block in build_outline(data, mode=mode):
            if block['kind'] == 'group' and block['transcript']:
                story += [Spacer(1, 10), para(f"Audio skripti{': ' + block['title'] if block['title'] else ''}", h2),
                          para(block['transcript'], body)]
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=_title(data)).build(story)
    return buffer.getvalue()


# ─── Excel / CSV / JSON ───────────────────────────────────────────────────────


def _rows(data: dict, mode: str) -> list:
    teacher = mode == 'teacher'
    head = ['№', 'Matn parchasi', 'Savol', 'Tur', 'Ball', 'Variantlar']
    if teacher:
        head.append("To'g'ri javob")
    rows = [head]
    group_title = ''
    for block in build_outline(data, mode=mode):
        if block['kind'] == 'group':
            group_title = block['title'] or ('Audio' if block['audio'] else '')
            continue
        variants = '\n'.join(text for text, _ in block['lines'])
        row = [block['n'], group_title, block['text'], block['type'], block['points'] or '', variants]
        if teacher:
            row.append(block['answer'])
        rows.append(row)
    return rows


def to_xlsx(data: dict, mode: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Test'[:31]
    for row in _rows(data, mode):
        sheet.append([_safe_cell(v) for v in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip('ABCDEFG', (5, 22, 60, 12, 7, 50, 28)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _safe_cell(value):
    """Excel formula sifatida o'qiladigan qiymatlar (=, +, -, @) matn bo'lib qolsin (formula inyeksiyasidan himoya)."""
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def to_csv(data: dict, mode: str) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for row in _rows(data, mode):
        writer.writerow([_safe_cell(v) for v in row])
    return ('﻿' + buffer.getvalue()).encode('utf-8')  # BOM: Excel o'zbekcha harflarni to'g'ri ochadi


def to_json(data: dict, mode: str) -> bytes:
    payload = {'title': _title(data), 'mode': mode, 'subject': data.get('subject_label') or '', 'questions': []}
    for block in build_outline(data, mode=mode):
        if block['kind'] == 'group':
            payload.setdefault('groups', []).append({
                'title': block['title'], 'passage': block['passage'], 'audio': block['audio'],
                **({'transcript': block['transcript']} if mode == 'teacher' else {}),
            })
            continue
        item = {'number': block['n'], 'text': block['text'], 'type': block['type'], 'points': block['points'],
                'options': [text for text, _ in block['lines']]}
        if mode == 'teacher':
            item['answer'] = block['answer']
        payload['questions'].append(item)
    return json.dumps(payload, ensure_ascii=False, indent=2).encode('utf-8')


_BUILDERS = {'pdf': to_pdf, 'docx': to_docx, 'xlsx': to_xlsx, 'csv': to_csv, 'json': to_json}


def export_quiz(data: dict, *, fmt: str, mode: str) -> tuple:
    """(bayt, content-type) qaytaradi. `fmt`/`mode` noto'g'ri bo'lsa ValueError."""
    if fmt not in _BUILDERS:
        raise ValueError(f"Format noto'g'ri: {', '.join(FORMATS)} dan birini tanlang.")
    if mode not in MODES:
        raise ValueError("Rejim noto'g'ri: student yoki teacher.")
    return _BUILDERS[fmt](data, mode), FORMATS[fmt]
