"""Testni PDF / Word / Excel / CSV / JSON qilib yuklab olish: javob kaliti faqat o'qituvchi nusxasida bo'lishi shart."""
import io
import json
import shutil
import tempfile
import zipfile

from django.core.files.base import ContentFile
from django.test import override_settings

from . import export, services
from .models import Quiz
from .tests_question_types import QuizTestBase

SECRET_FILL = 'Thompson'
SECRET_NUMBER = '4242'
SECRET_TEXT = 'photosynthesis'
SECRET_MATCH_RIGHT = 'Capital of France'
TRANSCRIPT = 'A: Good morning.\nB: My name is Thompson, T H O M P S O N.'

QUESTIONS = [
    {'text': 'Poytaxt qaysi?', 'type': 'single', 'points': 1, 'group': 0, 'options': [
        {'text': "Toshkent", 'is_correct': True, 'order': 0}, {'text': 'Samarqand', 'is_correct': False, 'order': 1}]},
    {'text': 'Tub sonlarni tanlang', 'type': 'multiple', 'points': 2, 'group': 0, 'options': [
        {'text': '2', 'is_correct': True, 'order': 0}, {'text': '4', 'is_correct': False, 'order': 1},
        {'text': '7', 'is_correct': True, 'order': 2}]},
    {'text': 'Yer yumaloq.', 'type': 'true_false', 'points': 1, 'correct_bool': True},
    {'text': 'Moslang', 'type': 'matching', 'points': 2,
     'pairs': [{'left': 'Paris', 'right': SECRET_MATCH_RIGHT}, {'left': 'Rome', 'right': 'Capital of Italy'}]},
    {'text': 'Tartibga soling', 'type': 'ordering', 'points': 2, 'items': ['birinchi', 'ikkinchi', 'uchinchi']},
    {'text': f'Surname: {{{{1}}}}, phone: {{{{2}}}}', 'type': 'fill_blank', 'points': 2,
     'blanks': [{'answers': [SECRET_FILL]}, {'answers': ['0771']}]},
    {'text': 'Jarayon nomi?', 'type': 'text', 'points': 1, 'accepted_answers': [SECRET_TEXT], 'case_sensitive': False},
    {'text': '2 x 2121 = ?', 'type': 'numeric', 'points': 1, 'accepted_answers': [SECRET_NUMBER], 'tolerance': 0},
]


class ExportBase(QuizTestBase):
    def setUp(self):
        super().setUp()
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)
        teacher = self.course_teacher()
        self.quiz = services.create_quiz(
            teacher=teacher, topic='Aralash test', title="O'zbekiston · oʻqish", course=self.course_obj(),
            questions=QUESTIONS, status=Quiz.Status.PUBLISHED,
            groups=[{'title': 'Listening 1', 'passage': 'Matn parchasi', 'transcript': TRANSCRIPT}],
        )
        group = self.quiz.groups.get()
        group.audio.save('a.mp3', ContentFile(b'ID3'), save=True)

    def course_obj(self):
        from apps.lessons.models import Course

        return Course.objects.get(pk=self.course_id)

    def course_teacher(self):
        return self.course_obj().teacher

    def get(self, fmt, mode=None, token=None):
        self.auth(token or self.teacher_token)
        url = f'/api/v1/quizzes/{self.quiz.id}/export/?type={fmt}'
        return self.client.get(url + (f'&mode={mode}' if mode else ''))

    @staticmethod
    def text_of(fmt, content):
        if fmt == 'docx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                return archive.read('word/document.xml').decode('utf-8')
        if fmt == 'xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                parts = [archive.read(n).decode('utf-8') for n in archive.namelist() if n.endswith('.xml')]
            return '\n'.join(parts)
        if fmt == 'pdf':
            from pypdf import PdfReader

            return '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(content)).pages)
        return content.decode('utf-8-sig')


class FormatTests(ExportBase):
    def test_every_format_downloads_with_the_right_type(self):
        types = {'pdf': 'application/pdf', 'docx': 'wordprocessingml', 'xlsx': 'spreadsheetml',
                 'csv': 'text/csv', 'json': 'application/json'}
        for fmt, expected in types.items():
            response = self.get(fmt, 'teacher')
            self.assertEqual(response.status_code, 200, (fmt, response.content[:200]))
            self.assertIn(expected, response['Content-Type'])
            self.assertIn(f'.{fmt}"', response['Content-Disposition'])
            self.assertGreater(len(response.content), 100, fmt)
        self.assertTrue(self.get('pdf', 'teacher').content.startswith(b'%PDF'))

    def test_default_is_a_student_pdf(self):
        response = self.get('pdf')
        self.assertEqual(response.status_code, 200)
        self.assertIn('_student.pdf', response['Content-Disposition'])

    def test_the_excel_file_opens_and_has_a_row_per_question(self):
        import openpyxl

        content = self.get('xlsx', 'teacher').content
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        self.assertEqual(sheet.max_row, 1 + len(QUESTIONS))
        self.assertEqual(sheet['G1'].value, "To'g'ri javob")

    def test_the_word_file_is_a_real_docx(self):
        import docx

        document = docx.Document(io.BytesIO(self.get('docx', 'teacher').content))
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('Poytaxt qaysi?', text)
        self.assertIn('Listening 1', text)

    def test_uzbek_apostrophe_letters_survive(self):
        text = self.text_of('pdf', self.get('pdf', 'teacher').content)
        self.assertIn('Poytaxt', text)
        self.assertIn('Tub sonlarni', text)

    def test_bad_type_or_mode_is_a_400(self):
        self.assertEqual(self.get('exe').status_code, 400)
        self.assertEqual(self.get('pdf', 'admin').status_code, 400)


class AnswerKeyTests(ExportBase):
    SECRETS = (SECRET_FILL, SECRET_NUMBER, SECRET_TEXT)

    def test_the_student_copy_never_contains_the_answer_key(self):
        for fmt in ('pdf', 'docx', 'xlsx', 'csv', 'json'):
            text = self.text_of(fmt, self.get(fmt, 'student').content)
            for secret in self.SECRETS:
                self.assertNotIn(secret, text, (fmt, secret))
            self.assertNotIn('✓', text, fmt)
            self.assertNotIn('Good morning', text, fmt)  # Listening skripti o'quvchiga berilmaydi
            self.assertIn('Poytaxt qaysi?', text, fmt)

    def test_the_teacher_copy_has_answers_and_the_listening_script(self):
        for fmt in ('pdf', 'docx', 'xlsx', 'csv', 'json'):
            text = self.text_of(fmt, self.get(fmt, 'teacher').content)
            for secret in self.SECRETS:
                self.assertIn(secret, text, (fmt, secret))
            if fmt not in ('xlsx', 'csv'):  # jadval: bir qator = bir savol, skriptga joy yo'q
                self.assertIn('Good morning', text, fmt)
        for fmt in ('pdf', 'docx'):
            self.assertIn('✓', self.text_of(fmt, self.get(fmt, 'teacher').content), fmt)

    def test_the_matching_answer_key_pairs_each_item(self):
        blocks = export.build_outline(self.get_data(), mode='teacher')
        matching = next(b for b in blocks if b['kind'] == 'question' and b['type'] == 'matching')
        self.assertRegex(matching['answer'], r'1-[A-B]; 2-[A-B]')
        student = next(b for b in export.build_outline(self.get_data(), mode='student')
                       if b['kind'] == 'question' and b['type'] == 'matching')
        self.assertEqual(student['answer'], '')

    def get_data(self):
        from rest_framework.test import APIRequestFactory

        from .serializers import QuizDetailSerializer

        return dict(QuizDetailSerializer(self.quiz, context={'request': APIRequestFactory().get('/')}).data)

    def test_fill_blank_gaps_are_numbered_for_the_student(self):
        question = next(b for b in export.build_outline(self.get_data(), mode='student')
                        if b['kind'] == 'question' and b['type'] == 'fill_blank')
        self.assertIn('______(1)', question['text'])
        self.assertIn('______(2)', question['text'])

    def test_the_ordering_items_are_shuffled_the_same_way_every_time(self):
        first = [b['lines'] for b in export.build_outline(self.get_data(), mode='student') if b.get('type') == 'ordering']
        second = [b['lines'] for b in export.build_outline(self.get_data(), mode='student') if b.get('type') == 'ordering']
        self.assertEqual(first, second)


class SafetyTests(ExportBase):
    def test_students_and_strangers_cannot_export(self):
        self.assertEqual(self.get('pdf', 'teacher', token=self.child_token).status_code, 403)
        self.assertEqual(self.get('pdf', 'student', token=self.parent_token).status_code, 403)

    def test_another_teachers_quiz_is_not_found(self):
        from apps.accounts.models import User
        from rest_framework_simplejwt.tokens import RefreshToken

        other = User.objects.create_user(username='other_t', password='x', role=User.Role.TEACHER, is_approved=True)
        token = str(RefreshToken.for_user(other).access_token)
        self.assertEqual(self.get('pdf', 'teacher', token=token).status_code, 404)

    def test_excel_formulas_are_neutralised(self):
        rows = export._rows({'questions': [{'type': 'text', 'text': '=1+1', 'order': 0, 'accepted_answers': ['@x']}],
                             'groups': []}, 'teacher')
        self.assertEqual(rows[1][2], '=1+1')  # xom qator o'zgarmaydi...
        self.assertEqual(export._safe_cell(rows[1][2]), "'=1+1")  # ...Excel/CSV ga yozishda esa neytrallanadi
        self.assertEqual(export._safe_cell('oddiy'), 'oddiy')

    def test_the_plain_json_is_valid(self):
        payload = json.loads(self.get('json', 'student').content)
        self.assertEqual(len(payload['questions']), len(QUESTIONS))
        self.assertNotIn('answer', payload['questions'][0])
